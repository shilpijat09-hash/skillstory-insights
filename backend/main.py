
import os
import json
import logging
from datetime import datetime, timezone
from typing import Any

import httpx
from dotenv import load_dotenv
from fastapi import BackgroundTasks, FastAPI, Header, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from supabase import create_client, Client
from pypdf import PdfReader
from docx import Document
from io import BytesIO

from crewai import Agent, Task, Crew, LLM
from backend.google_sheets import append_completed_analysis

load_dotenv("backend/.env")

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("resume-reviewer")

SUPABASE_URL = os.getenv("SUPABASE_URL", "").rstrip("/")
SUPABASE_SERVICE_ROLE_KEY = os.getenv("SUPABASE_SERVICE_ROLE_KEY", "")
OPENAI_API_KEY = os.getenv("OPENAI_API_KEY", "")

app = FastAPI(title="Resume Reviewer API", version="1.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:5173",
        "http://127.0.0.1:5173",
        "http://localhost:8080",
        "http://127.0.0.1:8080",
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

class AnalysisRequest(BaseModel):
    analysis_id: str


def get_admin_client() -> Client:
    if not SUPABASE_URL or not SUPABASE_SERVICE_ROLE_KEY:
        raise RuntimeError("Backend Supabase configuration is missing.")
    return create_client(SUPABASE_URL, SUPABASE_SERVICE_ROLE_KEY)


async def get_authenticated_user_id(
    authorization: str | None,
) -> str:
    if not authorization or not authorization.lower().startswith("bearer "):
        raise HTTPException(status_code=401, detail="Bearer token required.")

    token = authorization.split(" ", 1)[1].strip()
    if not token:
        raise HTTPException(status_code=401, detail="Access token is empty.")

    if not SUPABASE_URL or not SUPABASE_SERVICE_ROLE_KEY:
        raise HTTPException(status_code=500, detail="Backend configuration is missing.")

    try:
        async with httpx.AsyncClient(timeout=15) as client:
            response = await client.get(
                f"{SUPABASE_URL}/auth/v1/user",
                headers={
                    "apikey": SUPABASE_SERVICE_ROLE_KEY,
                    "Authorization": f"Bearer {token}",
                },
            )
    except httpx.HTTPError:
        raise HTTPException(status_code=503, detail="Could not contact Supabase Auth.")

    if response.status_code != 200:
        raise HTTPException(status_code=401, detail="Invalid or expired access token.")

    user_id = response.json().get("id")
    if not user_id:
        raise HTTPException(status_code=401, detail="Could not identify user.")

    return user_id


@app.get("/health")
async def health() -> dict[str, str]:
    return {"status": "ok", "service": "resume-reviewer-api"}


def extract_resume_text(file_bytes: bytes, file_name: str, file_type: str) -> str:
    name = file_name.lower()
    kind = (file_type or "").lower()

    if name.endswith(".pdf") or "pdf" in kind:
        reader = PdfReader(BytesIO(file_bytes))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif name.endswith(".docx") or "wordprocessingml" in kind:
        document = Document(BytesIO(file_bytes))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        for table in document.tables:
            for row in table.rows:
                text += "\n" + " | ".join(cell.text for cell in row.cells)
    else:
        raise ValueError("Unsupported resume format. Please upload a PDF or DOCX.")

    text = text.strip()
    if len(text) < 50:
        raise ValueError(
            "Could not extract enough text from this resume. "
            "Scanned PDFs may need OCR before analysis."
        )

    return text[:30000]


def parse_report(raw: str) -> dict[str, Any]:
    text = raw.strip()

    if text.startswith("```"):
        lines = text.splitlines()
        if lines and lines[0].startswith("```"):
            lines = lines[1:]
        if lines and lines[-1].strip() == "```":
            lines = lines[:-1]
        text = "\n".join(lines).strip()

    start = text.find("{")
    end = text.rfind("}")
    if start < 0 or end <= start:
        raise ValueError("The AI did not return a valid JSON report.")

    report = json.loads(text[start:end + 1])

    required = [
        "summary", "overallScore", "scoreBreakdown", "skills",
        "strengths", "weaknesses", "missingKeywords", "feedback",
    ]
    missing = [key for key in required if key not in report]
    if missing:
        raise ValueError("AI report is missing fields: " + ", ".join(missing))

    report["overallScore"] = max(0, min(100, int(report["overallScore"])))
    report["meta"] = {
        "engine": "crewai",
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "agents": ["Resume Analyzer", "Career Coach"],
        "model": "gpt-4o-mini",
    }
    return report


def run_analysis(analysis_id: str, user_id: str) -> None:
    supabase = get_admin_client()

    try:
        result = (
            supabase.table("analyses")
            .select("*")
            .eq("id", analysis_id)
            .eq("user_id", user_id)
            .limit(1)
            .execute()
        )
        if not result.data:
            logger.warning("Analysis %s was not found for its owner.", analysis_id)
            return

        analysis = result.data[0]
        if analysis.get("status") == "completed":
            return

        supabase.table("analyses").update({
            "status": "processing",
            "error_message": None,
        }).eq("id", analysis_id).eq("user_id", user_id).execute()

        if not OPENAI_API_KEY:
            raise RuntimeError("OPENAI_API_KEY is not configured.")

        resume_path = analysis.get("resume_path")
        if not resume_path:
            raise ValueError("No resume file path is saved for this analysis.")

        # The private Storage bucket is named "resumes".
        file_bytes = supabase.storage.from_("resumes").download(resume_path)

        resume_text = extract_resume_text(
            file_bytes,
            analysis.get("file_name", ""),
            analysis.get("file_type", ""),
        )

        job_description = (analysis.get("job_description") or "").strip()
        if not job_description:
            raise ValueError("The job description is empty.")

        llm = LLM(
            model="openai/gpt-4o-mini",
            api_key=OPENAI_API_KEY,
            temperature=0.2,
        )

        analyzer = Agent(
            role="Resume Analyzer",
            goal="Evaluate a resume against a specific job description using evidence.",
            backstory=(
                "You are an experienced ATS and recruitment analyst. "
                "You never invent qualifications that are absent from the resume."
            ),
            llm=llm,
            verbose=False,
        )

        coach = Agent(
            role="Career Coach",
            goal="Turn the evidence-based evaluation into practical resume improvements.",
            backstory=(
                "You give specific, ethical career advice. Suggestions must not "
                "claim the candidate has experience or skills they do not have."
            ),
            llm=llm,
            verbose=False,
        )

        task = Task(
            description=f"""
Analyze the resume against the job description below.

JOB TITLE: {analysis.get("job_title", "Not specified")}
COMPANY: {analysis.get("company") or "Not specified"}

JOB DESCRIPTION:
{job_description[:12000]}

RESUME TEXT:
{resume_text}

Return ONLY one valid JSON object, without markdown fences, with this exact structure:
{{
  "summary": "A concise evidence-based summary",
  "overallScore": 75,
  "scoreBreakdown": [
    {{"key":"ats","label":"ATS Compatibility","score":75,"description":"..."}},
    {{"key":"keywords","label":"Keyword Match","score":75,"description":"..."}},
    {{"key":"impact","label":"Impact","score":75,"description":"..."}},
    {{"key":"formatting","label":"Formatting","score":75,"description":"..."}},
    {{"key":"clarity","label":"Clarity","score":75,"description":"..."}}
  ],
  "skills": [
    {{"name":"Python","category":"technical","matchesJob":true}}
  ],
  "strengths": [
    {{"title":"Relevant experience","detail":"Evidence from the resume"}}
  ],
  "weaknesses": [
    {{"title":"Missing evidence","detail":"Explain the gap honestly"}}
  ],
  "missingKeywords": [
    {{"keyword":"Example keyword","importance":"high","suggestion":"How to address it truthfully"}}
  ],
  "feedback": [
    {{
      "id":"feedback-1",
      "section":"Experience",
      "priority":"high",
      "title":"Make an achievement more specific",
      "detail":"Explain the improvement",
      "example":{{"before":"Original-style example","after":"Suggested truthful rewrite"}}
    }}
  ]
}}

Rules:
- Scores must be integers from 0 to 100.
- Score dimension keys must be ats, keywords, impact, formatting, or clarity.
- Skill categories must be technical, tools, soft, or domain.
- Priorities must be high, medium, or low.
- Use empty arrays when there is no reliable evidence.
- Do not invent resume facts, metrics, degrees, or experience.
- Clearly distinguish missing evidence from proof that a skill is absent.
- Give useful feedback specific to this resume and job.
- Keep the report concise and valid JSON.
""",
            expected_output="One valid JSON object matching the requested report structure.",
            agent=analyzer,
        )

        review_task = Task(
            description=(
                "Review the previous task's proposed report for factual restraint, "
                "usefulness, valid enum values and valid JSON. Preserve the exact "
                "required schema. Return only the corrected JSON object."
            ),
            expected_output="A corrected, valid JSON report with the required fields.",
            agent=coach,
            context=[task],
        )

        crew = Crew(
            agents=[analyzer, coach],
            tasks=[task, review_task],
            verbose=False,
        )

        result = crew.kickoff()
        report = parse_report(result.raw)

        completed_at = datetime.now(timezone.utc).isoformat()
        supabase.table("analyses").update({
            "status": "completed",
            "engine": "crewai",
            "score": report["overallScore"],
            "report": report,
            "error_message": None,
            "completed_at": completed_at,
        }).eq("id", analysis_id).eq("user_id", user_id).execute()

        logger.info("Analysis %s completed.", analysis_id)
        append_completed_analysis({
            "analysis_id": analysis_id,
            "user_id": user_id,
            "job_title": analysis.get("job_title", ""),
            "company": analysis.get("company"),
            "overall_score": report["overallScore"],
            "status": "completed",
            "completed_at": completed_at,
        })

    except Exception as exc:
        logger.exception("Analysis %s failed.", analysis_id)
        try:
            supabase.table("analyses").update({
                "status": "failed",
                "error_message": str(exc)[:1000],
            }).eq("id", analysis_id).eq("user_id", user_id).execute()
        except Exception:
            logger.exception("Could not save failure status for %s.", analysis_id)


@app.post("/v1/analyses")
async def start_analysis(
    body: AnalysisRequest,
    background_tasks: BackgroundTasks,
    authorization: str | None = Header(default=None),
) -> dict[str, Any]:
    user_id = await get_authenticated_user_id(authorization)
    supabase = get_admin_client()

    result = (
        supabase.table("analyses")
        .select("*")
        .eq("id", body.analysis_id)
        .eq("user_id", user_id)
        .limit(1)
        .execute()
    )

    if not result.data:
        raise HTTPException(status_code=404, detail="Analysis not found.")

    analysis = result.data[0]
    status = analysis.get("status")

    if status == "completed":
        return {"status": "completed", "analysis_id": body.analysis_id}

    if status == "processing":
        return {"status": "processing", "analysis_id": body.analysis_id}

    if status not in ("queued", "failed"):
        raise HTTPException(
            status_code=409,
            detail="Analysis cannot be started in its current state.",
        )

    if not analysis.get("resume_path"):
        raise HTTPException(status_code=400, detail="Resume path is missing.")

    supabase.table("analyses").update({
        "status": "processing",
        "error_message": None,
    }).eq("id", body.analysis_id).eq("user_id", user_id).execute()

    background_tasks.add_task(run_analysis, body.analysis_id, user_id)

    return {"status": "processing", "analysis_id": body.analysis_id}
